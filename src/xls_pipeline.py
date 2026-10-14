# xls_pipeline.py
#
# Write categories to an xlsx file.

import sys, os, json
import glob
from os.path import abspath, basename, join, dirname
from openpyxl import Workbook
from string import ascii_uppercase as letters

ROOT = dirname(__file__)
SECTIONS = ('e', 'm', 'r', 's')

letters = [l for l in letters] + [i+i for i in letters]
# print("\n", letters)

files = glob.glob(abspath(join(ROOT, "categories", "*-*.json")))

wb = Workbook()
ws = wb.active

l = 0
for i, f in enumerate(files):
    print(basename(f))

    with open(f, 'r') as file:
        dct = json.load(file)

    # print(dct['test_code'])

    col = letters[l]
    row = 1
    ws[col + str(row)] = dct['test_code']

    # e, m, r, s sections
    for code in SECTIONS:
        questions = dct[code]['q']
        row += 1
        ws[col + str(row)] = code

        # grab cat list from question dict
        for q, d in questions.items():
            cats = ", ".join(d['cat'])
            row += 1
            ws[col + str(row)] = cats
        
    

    # jump the letter/col index for the next iteration
    l += 2

wb.save(abspath(join(ROOT, "categories", "Official ACT Categories.xlsx")))
    












